"""Regenerate synthetic extraction fixtures with the bundled artifact runtime."""
from pathlib import Path
from reportlab.pdfgen import canvas
from reportlab.lib.pagesizes import letter
from docx import Document
from docx.shared import Inches, Pt

target = Path(__file__).resolve().parents[1] / "tests" / "fixtures" / "screening"
target.mkdir(parents=True, exist_ok=True)
lines = [
    "Ada Synthetic",
    "Backend Engineer",
    "Experience",
    "Built TypeScript APIs at Example Studio from 2020 to 2024.",
    "Skills",
    "TypeScript, PostgreSQL, Node.js",
    "Education",
    "BSc Computer Science, Example University, 2019.",
]
pdf = canvas.Canvas(str(target / "synthetic-resume.pdf"), pagesize=letter)
pdf.setTitle("Synthetic resume extraction fixture")
pdf.setAuthor("Recruta test fixtures")
for index, line in enumerate(lines):
    pdf.setFont("Helvetica-Bold" if index in (0, 2, 4, 6) else "Helvetica", 16 if index == 0 else 11)
    pdf.drawString(54, 738 - index * 28, line)
pdf.save()

doc = Document()
section = doc.sections[0]
section.page_width, section.page_height = Inches(8.5), Inches(11)
section.top_margin = section.bottom_margin = Inches(0.75)
doc.styles["Normal"].font.name = "Arial"
doc.styles["Normal"].font.size = Pt(11)
doc.add_paragraph(lines[0], style="Title")
for index, line in enumerate(lines[1:], start=1):
    doc.add_paragraph(line, style="Heading 1" if index in (2, 4, 6) else "Normal")
doc.core_properties.author = "Recruta test fixtures"
doc.core_properties.title = "Synthetic resume extraction fixture"
doc.save(target / "synthetic-resume.docx")
print("Created two synthetic extraction fixtures.")
